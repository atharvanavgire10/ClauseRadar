"""Phase 03 tests — document ingestion pipeline (PDF/DOCX, states, isolation)."""
import io

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from rest_framework.test import APIClient

from audit.models import AuditEvent
from contracts.models import Contract
from documents.models import Document
from organizations.models import Organization, OrganizationMembership
from workspaces.models import Workspace, WorkspaceMembership

User = get_user_model()


def make_pdf_bytes(pages: int = 1, prefix: str = "ClauseRadar test page") -> bytes:
    import fitz

    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"{prefix} {i + 1}. The Vendor shall maintain insurance.")
    data = doc.tobytes()
    doc.close()
    return bytes(data)


def make_docx_bytes(paragraphs: int = 5, text: str = "The Vendor shall deliver monthly reports.") -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    for i in range(paragraphs):
        doc.add_paragraph(f"{text} (para {i + 1})")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_empty_docx_bytes() -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def setup(db):
    user = User.objects.create_user(email="owner@example.com", password="password123")
    other = User.objects.create_user(email="other@example.com", password="password123")
    org = Organization.objects.create(name="Acme", created_by=user)
    OrganizationMembership.objects.create(organization=org, user=user, role="OWNER")
    ws = Workspace.objects.create(organization=org, name="Legal", created_by=user)
    WorkspaceMembership.objects.create(workspace=ws, user=user, role="OWNER")
    contract = Contract.objects.create(workspace=ws, title="Vendor Agreement", created_by=user)
    return {"user": user, "other": other, "org": org, "ws": ws, "contract": contract}


def auth(user):
    client = APIClient()
    client.force_authenticate(user=user)
    return client


def upload(client, contract_id, filename, content, content_type="application/octet-stream"):
    f = SimpleUploadedFile(filename, content, content_type=content_type)
    return client.post("/api/v1/documents/", {"contract": str(contract_id), "file": f}, format="multipart")


def test_pdf_upload_ready_with_pages(setup):
    client = auth(setup["user"])
    r = upload(client, setup["contract"].id, "agreement.pdf", make_pdf_bytes(2), "application/pdf")
    assert r.status_code == 201, r.content
    body = r.json()
    assert body["status"] == "READY"
    assert body["page_count"] == 2
    assert body["mime"] == "application/pdf"

    doc_id = body["id"]
    pages = client.get(f"/api/v1/documents/{doc_id}/pages/").json()
    results = pages["results"] if "results" in pages else pages
    assert len(results) == 2
    assert "Vendor shall maintain insurance" in results[0]["text"]
    assert AuditEvent.objects.filter(entity_type="document", action="document.processed").exists()


def test_docx_upload_ready(setup):
    client = auth(setup["user"])
    r = upload(client, setup["contract"].id, " sow.docx ", make_docx_bytes(3), "application/octet-stream")
    assert r.status_code == 201, r.content
    body = r.json()
    assert body["status"] == "READY"
    assert body["original_filename"] == "sow.docx"  # sanitized, never trust raw name
    assert body["page_count"] >= 1


def test_duplicate_upload_returns_409(setup):
    client = auth(setup["user"])
    pdf = make_pdf_bytes(1)
    r1 = upload(client, setup["contract"].id, "a.pdf", pdf, "application/pdf")
    assert r1.status_code == 201
    r2 = upload(client, setup["contract"].id, "renamed-copy.pdf", pdf, "application/pdf")
    assert r2.status_code == 409
    assert r2.json()["code"] == "duplicate"
    assert r2.json()["existing_id"] == r1.json()["id"]


def test_renamed_executable_rejected(setup):
    client = auth(setup["user"])
    r = upload(client, setup["contract"].id, "malware.pdf", b"MZ\x90\x00evil binary", "application/pdf")
    assert r.status_code == 400
    assert r.json()["code"] == "unsupported_type"


def test_txt_file_rejected(setup):
    client = auth(setup["user"])
    r = upload(client, setup["contract"].id, "notes.txt", b"just text", "text/plain")
    assert r.status_code == 400


def test_corrupt_pdf_marked_failed_not_crash(setup):
    client = auth(setup["user"])
    # Passes magic sniff (%PDF) but is not a parseable PDF → FAILED state, not 500.
    r = upload(client, setup["contract"].id, "broken.pdf", b"%PDF-1.4 not really a pdf \x00\x01\x02", "application/pdf")
    assert r.status_code == 201, r.content
    assert r.json()["status"] == "FAILED"
    assert r.json()["error_message"]
    assert AuditEvent.objects.filter(action="document.processing_failed").exists()


def test_empty_docx_marked_failed(setup):
    client = auth(setup["user"])
    r = upload(client, setup["contract"].id, "empty.docx", make_empty_docx_bytes())
    assert r.status_code == 201, r.content
    assert r.json()["status"] == "FAILED"
    assert "no readable text" in r.json()["error_message"].lower()


@override_settings(MAX_UPLOAD_MB=0)
def test_size_limit_enforced(setup):
    client = auth(setup["user"])
    r = upload(client, setup["contract"].id, "big.pdf", make_pdf_bytes(1))
    assert r.status_code == 400
    assert r.json()["code"] == "file_too_large"


def test_cross_workspace_isolation(setup):
    owner_client = auth(setup["user"])
    other_client = auth(setup["other"])
    r = upload(owner_client, setup["contract"].id, "secret.pdf", make_pdf_bytes(1))
    doc_id = r.json()["id"]
    assert other_client.get("/api/v1/documents/").json()["count"] == 0
    assert other_client.get(f"/api/v1/documents/{doc_id}/").status_code == 404
    assert other_client.get(f"/api/v1/documents/{doc_id}/download/").status_code == 404


def test_download_returns_source_file(setup):
    client = auth(setup["user"])
    pdf = make_pdf_bytes(1)
    doc_id = upload(client, setup["contract"].id, "dl.pdf", pdf).json()["id"]
    r = client.get(f"/api/v1/documents/{doc_id}/download/")
    assert r.status_code == 200
    assert r["Content-Type"] == "application/pdf"
    assert b"%PDF" in b"".join(r.streaming_content)


def test_reprocess_is_idempotent(setup):
    client = auth(setup["user"])
    doc_id = upload(client, setup["contract"].id, "re.pdf", make_pdf_bytes(2)).json()["id"]
    r1 = client.post(f"/api/v1/documents/{doc_id}/reprocess/")
    assert r1.status_code == 200
    assert r1.json()["page_count"] == 2
    pages = client.get(f"/api/v1/documents/{doc_id}/pages/").json()
    results = pages["results"] if "results" in pages else pages
    assert len(results) == 2  # no duplicates after reprocess


def test_viewer_cannot_upload(setup):
    org2 = Organization.objects.create(name="Other", created_by=setup["other"])
    OrganizationMembership.objects.create(organization=org2, user=setup["other"], role="MEMBER")
    WorkspaceMembership.objects.create(workspace=setup["ws"], user=setup["other"], role="VIEWER")
    client = auth(setup["other"])
    r = upload(client, setup["contract"].id, "nope.pdf", make_pdf_bytes(1))
    assert r.status_code in (400, 403, 404)


def test_delete_removes_file_and_audits(setup):
    client = auth(setup["user"])
    doc_id = upload(client, setup["contract"].id, "todelete.pdf", make_pdf_bytes(1)).json()["id"]
    doc = Document.objects.get(pk=doc_id)
    stored = doc.file.name
    from django.core.files.storage import default_storage

    assert default_storage.exists(stored)
    assert client.delete(f"/api/v1/documents/{doc_id}/").status_code == 204
    assert not default_storage.exists(stored)
    assert AuditEvent.objects.filter(action="document.deleted").exists()
