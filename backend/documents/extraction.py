"""Upload validation + text extraction (PyMuPDF / python-docx / OCR fallback)."""
from __future__ import annotations

import hashlib
import io
import os
from pathlib import Path

ALLOWED_MIMES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
ALLOWED_EXTENSIONS = {".pdf", ".docx"}
MAGIC_PDF = b"%PDF"
MAGIC_ZIP = b"PK\x03\x04"

# Pseudo-page target for DOCX (which has no native pages).
DOCX_CHARS_PER_PAGE = 4000


def max_upload_bytes() -> int:
    from django.conf import settings

    return int(getattr(settings, "MAX_UPLOAD_MB", 25)) * 1024 * 1024


def sanitize_filename(name: str) -> str:
    base = os.path.basename(name or "upload").strip() or "upload"
    base = "".join(c if (c.isalnum() or c in "._- ") else "_" for c in base).strip(" .")
    return base[:200] or "upload"


def sniff_mime(filename: str, head: bytes, declared: str) -> str | None:
    """Return canonical mime if the file looks like PDF/DOCX, else None.

    Never trusts the client-supplied content type or extension alone: magic
    bytes must agree with the extension.
    """
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        return None
    if ext == ".pdf" and head.startswith(MAGIC_PDF):
        return "application/pdf"
    if ext == ".docx" and head.startswith(MAGIC_ZIP):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    # Declared content-type is only a hint; magic bytes decide.
    if declared in ALLOWED_MIMES and ext == (".pdf" if declared == "application/pdf" else ".docx"):
        # Magic mismatch → reject (prevents renamed executables).
        return None
    return None


def compute_sha256(chunks) -> str:
    h = hashlib.sha256()
    for chunk in chunks:
        h.update(chunk)
    return h.hexdigest()


def extract_pdf_pages(file_path: str) -> tuple[list[str], bool]:
    """Return (page_texts, ocr_used). Raises RuntimeError on unreadable PDFs."""
    import fitz  # PyMuPDF

    try:
        doc = fitz.open(file_path)
    except Exception as exc:
        raise RuntimeError(f"Could not open PDF: {exc}") from exc
    texts: list[str] = []
    ocr_used = False
    try:
        if len(doc) == 0:
            raise RuntimeError("PDF contains no pages.")
        for i, page in enumerate(doc):
            text = (page.get_text("text") or "").strip()
            if not text:
                ocr_text = _try_ocr_page(page)
                if ocr_text:
                    text = ocr_text
                    ocr_used = True
            texts.append(text)
    finally:
        doc.close()
    if not texts:
        raise RuntimeError("PDF contains no pages.")
    return texts, ocr_used


def _try_ocr_page(page) -> str:
    """Best-effort OCR for scanned pages. Returns '' when unavailable."""
    try:
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore
    except ImportError:
        return ""
    try:
        pix = page.get_pixmap(dpi=150)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        return (pytesseract.image_to_string(img) or "").strip()
    except Exception:
        return ""


def extract_docx_pages(file_path: str) -> tuple[list[str], bool]:
    """Split DOCX paragraphs into pseudo-pages. Raises RuntimeError when empty/unreadable."""
    from docx import Document as DocxDocument

    try:
        doc = DocxDocument(file_path)
    except Exception as exc:
        raise RuntimeError(f"Could not open DOCX: {exc}") from exc
    paragraphs = [(p.text or "").rstrip() for p in doc.paragraphs]
    # Include tables — obligations often live in tables.
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                paragraphs.append(line)
    paragraphs = [p for p in paragraphs if p.strip()]
    if not paragraphs:
        raise RuntimeError("DOCX contains no readable text.")
    pages: list[str] = []
    current: list[str] = []
    current_len = 0
    for para in paragraphs:
        current.append(para)
        current_len += len(para) + 1
        if current_len >= DOCX_CHARS_PER_PAGE:
            pages.append("\n".join(current).strip())
            current, current_len = [], 0
    if current:
        pages.append("\n".join(current).strip())
    return pages, False
